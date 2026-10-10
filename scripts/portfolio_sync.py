"""Read-only Graph workbook extraction and authenticated portfolio verification.

No workbook, holdings, token, URL, response body, or private identifiers are logged.
The default mode compares against production without writing to it.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from io import BytesIO
import json
import math
import os
from pathlib import Path
import random
import subprocess
import time
from urllib import error, parse, request
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

MAPPING = json.loads(Path(__file__).with_name('asset-tables.json').read_text())
MAX_BYTES = 50 * 1024 * 1024
MAX_UNPACKED_BYTES = 200 * 1024 * 1024


class SyncError(Exception):
    """Safe, fixed error messages suitable for scheduler logs."""


def iso(value):
    try:
        stamp = datetime.fromisoformat(value.replace('Z', '+00:00'))
        if stamp.tzinfo is None:
            raise ValueError()
        return stamp.astimezone(timezone.utc)
    except (ValueError, AttributeError, TypeError):
        raise SyncError('invalid_timestamp') from None


def utcnow():
    return datetime.now(timezone.utc).isoformat().replace('+00:00', 'Z')


def validate_assets(assets):
    if not isinstance(assets, dict) or set(assets) != set(MAPPING):
        raise SyncError('incorrect_asset_set')
    for value in assets.values():
        if isinstance(value, bool) or not isinstance(value, (float, int)):
            raise SyncError('quantity_must_be_numeric')
        if not math.isfinite(value) or value < 0:
            raise SyncError('invalid_quantity')
    return assets


def extract_quantities(content):
    if len(content) > MAX_BYTES:
        raise SyncError('workbook_too_large')
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(f.file_size for f in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise SyncError('workbook_expansion_limit')
        # XML is protected by installed defusedxml. Macros and formulas are never run.
        formulas = load_workbook(BytesIO(content), data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(content), data_only=True, keep_links=False)
    except SyncError:
        raise
    except Exception:
        raise SyncError('invalid_workbook') from None
    try:
        tables = {}
        for sheet in formulas:
            for table in sheet.tables.values():
                if table.name in tables:
                    raise SyncError('duplicate_table')
                tables[table.name] = (sheet, table)
        result = {}
        for asset, spec in MAPPING.items():
            entry = tables.get(spec['table'])
            if entry is None:
                raise SyncError('required_table_missing')
            sheet, table = entry
            left, top, right, bottom = range_boundaries(table.ref)
            columns = [c for c in table.tableColumns if c.name == spec['column']]
            if len(columns) != 1 or len(table.tableColumns) != right - left + 1:
                raise SyncError('quantity_column_ambiguous')
            column = columns[0]
            # Require an explicit Excel totals row. Never choose a transaction row.
            if table.totalsRowCount != 1 or table.totalsRowShown is False:
                raise SyncError('current_balance_total_not_explicit')
            if not column.totalsRowFunction and column.totalsRowFormula is None:
                raise SyncError('quantity_total_not_declared')
            index = left + list(table.tableColumns).index(column)
            if sheet.cell(top, index).value != spec['column']:
                raise SyncError('quantity_header_mismatch')
            if bottom <= top + 1:
                raise SyncError('invalid_table_range')
            value = cached[sheet.title].cell(bottom, index).value
            result[asset] = value
        return validate_assets(result)
    finally:
        formulas.close()
        cached.close()


def http(url, *, method='GET', headers=None, body=None, retries=3, redirect=True):
    """Bounded retries for safe/idempotent reads and OAuth requests only.

    Production POST callers set retries=0 because a timed-out write is ambiguous.
    """
    if parse.urlsplit(url).scheme != 'https':
        raise SyncError('https_required')
    class NoRedirect(request.HTTPRedirectHandler):
        def redirect_request(self, *args):
            return None
    opener = request.build_opener() if redirect else request.build_opener(NoRedirect)
    for attempt in range(retries + 1):
        try:
            req = request.Request(url, method=method, headers=headers or {}, data=body)
            with opener.open(req, timeout=30) as response:
                content = response.read(MAX_BYTES + 1)
                if len(content) > MAX_BYTES:
                    raise SyncError('response_too_large')
                return content
        except error.HTTPError as exc:
            status = exc.code
            if status not in (429, 500, 502, 503, 504) or attempt == retries:
                raise SyncError(f'http_{status}') from None
            delay = min(30, 2 ** attempt + random.random())
            retry_after = exc.headers.get('Retry-After', '')
            if retry_after.isdigit():
                delay = min(60, int(retry_after))
            exc.close()
        except (error.URLError, TimeoutError, OSError):
            if attempt == retries:
                raise SyncError('network_failure') from None
            delay = min(30, 2 ** attempt + random.random())
        time.sleep(delay)


def json_request(url, **kwargs):
    kwargs.setdefault('redirect', False)
    try:
        value = json.loads(http(url, **kwargs))
        if not isinstance(value, dict):
            raise ValueError()
        return value
    except (ValueError, UnicodeError):
        raise SyncError('invalid_service_response') from None


def required(name):
    value = os.environ.get(name)
    if not value:
        raise SyncError(f'missing_{name}')
    return value


def graph_token():
    mode = os.environ.get('MS_GRAPH_AUTH_MODE', 'delegated')
    tenant = os.environ.get('MS_GRAPH_TENANT_ID', 'consumers')
    if not tenant or any(c not in 'abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-.' for c in tenant):
        raise SyncError('invalid_graph_tenant')
    fields = {'client_id': required('MS_GRAPH_CLIENT_ID')}
    if mode == 'delegated':
        fields.update(grant_type='refresh_token', refresh_token=required('MS_GRAPH_REFRESH_TOKEN'))
        if os.environ.get('MS_GRAPH_CLIENT_SECRET'):
            fields['client_secret'] = os.environ['MS_GRAPH_CLIENT_SECRET']
    elif mode == 'application':
        if tenant in ('consumers', 'common', 'organizations'):
            raise SyncError('application_auth_requires_tenant')
        fields.update(grant_type='client_credentials', scope='https://graph.microsoft.com/.default',
                      client_secret=required('MS_GRAPH_CLIENT_SECRET'))
    else:
        raise SyncError('unsupported_graph_auth_mode')
    token = json_request(f'https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token',
                         method='POST', headers={'Content-Type': 'application/x-www-form-urlencoded'},
                         body=parse.urlencode(fields).encode())
    if not token.get('access_token'):
        raise SyncError('graph_authentication_failed')
    if 'files.readwrite' in token.get('scope', '').lower():
        raise SyncError('unexpected_write_scope')
    rotated = token.get('refresh_token')
    if mode == 'delegated' and rotated:
        # Do not lose refreshed authorization between ephemeral Actions runners.
        # GH_TOKEN must be a dedicated repo Secrets-write credential, never GITHUB_TOKEN.
        if not os.environ.get('GH_TOKEN') or not os.environ.get('GITHUB_REPOSITORY'):
            raise SyncError('refresh_token_persistence_not_configured')
        outcome = subprocess.run(['gh', 'secret', 'set', 'MS_GRAPH_REFRESH_TOKEN', '--repo',
                                  os.environ['GITHUB_REPOSITORY']], input=rotated.encode(),
                                 stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if outcome.returncode:
            raise SyncError('refresh_token_persistence_failed')
    return token['access_token']


def download_workbook(token):
    drive = parse.quote(required('MS_GRAPH_DRIVE_ID'), safe='')
    item = parse.quote(required('MS_GRAPH_ITEM_ID'), safe='')
    url = f'https://graph.microsoft.com/v1.0/drives/{drive}/items/{item}'
    headers = {'Authorization': 'Bearer ' + token}
    metadata = json_request(url + '?$select=id,name,lastModifiedDateTime,eTag,file', headers=headers)
    if metadata.get('name') != 'Trading Journal.xlsm' or not metadata.get('file') or not metadata.get('eTag'):
        raise SyncError('workbook_identity_not_verified')
    iso(metadata.get('lastModifiedDateTime'))
    if iso(metadata['lastModifiedDateTime']) > datetime.now(timezone.utc):
        raise SyncError('future_workbook_timestamp')
    # /content redirects to a preauthorized file URL. Never forward Graph auth to it.
    # Fetch that URL from metadata without Authorization and never log it.
    download = json_request(url, headers=headers).get('@microsoft.graph.downloadUrl')
    if not download or parse.urlsplit(download).scheme != 'https':
        raise SyncError('workbook_download_url_unavailable')
    content = http(download)
    after = json_request(url + '?$select=eTag', headers=headers)
    if after.get('eTag') != metadata['eTag']:
        raise SyncError('workbook_changed_during_read')
    return content, metadata['lastModifiedDateTime']


def endpoint():
    value = required('PORTFOLIO_ENDPOINT')
    url = parse.urlsplit(value)
    if url.scheme != 'https' or not url.hostname or url.username or url.password or url.query or url.fragment:
        raise SyncError('invalid_portfolio_endpoint')
    return value


def read_snapshot(base, secret):
    response = json_request(base + '?data=portfolio', headers={'x-portfolio-sync-secret': secret}, redirect=False)
    if response.get('ok') is not True or not isinstance(response.get('snapshot'), dict):
        raise SyncError('snapshot_readback_invalid')
    snapshot = response['snapshot']
    validate_assets(snapshot.get('assets'))
    iso(snapshot.get('updated_at'))
    if snapshot.get('workbook_updated_at'):
        iso(snapshot['workbook_updated_at'])
    return snapshot


def verify_readback(snapshot, payload):
    if (snapshot.get('assets') != payload['assets'] or
            snapshot.get('updated_at') != payload['updated_at'] or
            snapshot.get('workbook_updated_at') != payload['workbook_updated_at']):
        raise SyncError('snapshot_readback_mismatch')


def synchronize(assets, workbook_stamp, *, write=False):
    validate_assets(assets)
    iso(workbook_stamp)
    if iso(workbook_stamp) > datetime.now(timezone.utc):
        raise SyncError('future_workbook_timestamp')
    base = endpoint()
    secret = required('PORTFOLIO_SYNC_SECRET_AGENT')
    previous = read_snapshot(base, secret)  # Fail closed on inaccessible/invalid state.
    if iso(workbook_stamp) < iso(previous.get('workbook_updated_at') or previous['updated_at']):
        raise SyncError('older_workbook_rejected')
    payload = {'version': 'Microsoft Graph verified table totals', 'updated_at': utcnow(),
               'workbook_updated_at': workbook_stamp, 'assets': assets}
    if not write:
        return {'mode': 'compare', 'matched': previous['assets'] == assets,
                'verified_assets': len(assets), 'production_written': False}
    if os.environ.get('PORTFOLIO_MIGRATION_APPROVED') != 'true':
        raise SyncError('migration_not_approved')
    response = json_request(base + '?sync=portfolio', method='POST',
                            headers={'x-portfolio-sync-secret': secret, 'Content-Type': 'application/json'},
                            body=json.dumps(payload, allow_nan=False).encode(), retries=0, redirect=False)
    if response.get('ok') is not True or response.get('stored') is not True or response.get('received') != 23:
        raise SyncError('sync_not_confirmed')
    verify_readback(read_snapshot(base, secret), payload)
    return {'mode': 'write', 'verified_assets': 23, 'readback_verified': True, 'production_written': True}


def monitor():
    base = endpoint()
    health = json_request(base, redirect=False)
    if health.get('ok') is not True or health.get('service') != 'Portfolio Agent Telegram Bot':
        raise SyncError('health_check_failed')
    snapshot = read_snapshot(base, required('PORTFOLIO_SYNC_SECRET_AGENT'))
    now = datetime.now(timezone.utc)
    age = (now - iso(snapshot['updated_at'])).total_seconds()
    source_age = (now - iso(snapshot.get('workbook_updated_at') or snapshot['updated_at'])).total_seconds()
    if age < 0 or age > 24 * 3600 or source_age < 0:
        raise SyncError('snapshot_stale_or_future')
    # An unchanged file may be old while a new verified read is fresh. Do not
    # confuse workbook edit age with failed synchronization.
    return {'health_verified': True, 'snapshot_age_hours': round(age / 3600, 2),
            'workbook_edit_age_hours': round(source_age / 3600, 2), 'verified_assets': 23}


def inspect_telegram():
    # Read-only Bot API methods; no messages sent and no webhook mutation.
    token = required('TELEGRAM_BOT_TOKEN')
    base = 'https://api.telegram.org/bot' + token
    identity = json_request(base + '/getMe', redirect=False)
    webhook = json_request(base + '/getWebhookInfo', redirect=False)
    if identity.get('ok') is not True or identity.get('result', {}).get('username') != 'SaRoPortfolioAgentBot':
        raise SyncError('telegram_identity_mismatch')
    if webhook.get('ok') is not True or webhook.get('result', {}).get('url') != endpoint():
        raise SyncError('telegram_webhook_mismatch')
    info = webhook['result']
    # Surface backlog and error age; a historical last_error_date alone is not
    # proof of a current outage. Bot API success is not a user command E2E test.
    return {'telegram_identity_verified': True, 'webhook_url_verified': True,
            'pending_update_count': info.get('pending_update_count', 0),
            'last_error_date': info.get('last_error_date'), 'telegram_commands_verified': False}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--write', action='store_true')
    parser.add_argument('--monitor', action='store_true')
    parser.add_argument('--telegram', action='store_true')
    args = parser.parse_args()
    if sum([args.write, args.monitor, args.telegram]) > 1:
        parser.error('--write, --monitor and --telegram are mutually exclusive')
    try:
        if args.telegram:
            result = inspect_telegram()
        elif args.monitor:
            result = monitor()
        else:
            if args.write and os.environ.get('PORTFOLIO_MIGRATION_APPROVED') != 'true':
                raise SyncError('migration_not_approved')
            content, modified = download_workbook(graph_token())
            result = synchronize(extract_quantities(content), modified, write=args.write)
        print(json.dumps({'ok': True, **result}))
        return 0
    except SyncError as exc:
        print(json.dumps({'ok': False, 'error': str(exc)}))
        return 1
    except Exception:
        print(json.dumps({'ok': False, 'error': 'unexpected_failure'}))
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
