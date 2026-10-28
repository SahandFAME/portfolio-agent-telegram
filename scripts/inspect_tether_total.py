"""Read the requested total formula/comment; publish only encrypted diagnostics."""
import base64
from collections import defaultdict
from datetime import date,datetime
import hashlib
from io import BytesIO
import json
import os
import subprocess
import tempfile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from gold_history import norm,reconstruct
from verify_onedrive import MAPPING,SyncError,download_workbook,extract_quantities,graph_token,required


def inspect_dust_tables(book,cached):
    """Narrow, encrypted source review for confirmed dust totals and old tables."""
    selected={v['table']:a for a,v in MAPPING.items() if a in {'ADA','SOL','ONDO','Link'}}
    mapped={v['table'] for v in MAPPING.values()};result=[]
    def value(v):return v.isoformat() if isinstance(v,(date,datetime)) else v
    for sheet in book:
        for table in sheet.tables.values():
            columns=[c.name for c in table.tableColumns]
            if table.name not in selected and table.name in mapped:continue
            if not any(c in columns for c in ('Quantity','حجم خرید')) or not any(c in columns for c in ('Date','تاریخ')):continue
            left,top,right,bottom=range_boundaries(table.ref)
            qi=left+columns.index('Quantity' if 'Quantity' in columns else 'حجم خرید')
            di=left+columns.index('Date' if 'Date' in columns else 'تاریخ')
            total=sheet.cell(bottom,qi);values=cached[sheet.title]
            result.append({'asset':selected.get(table.name),'sheet':sheet.title,'table':table.name,'range':table.ref,
                'labels':{sheet.cell(r,c).coordinate:str(sheet.cell(r,c).value)[:100] for r in range(max(1,top-2),top+1) for c in range(left,right+1) if isinstance(sheet.cell(r,c).value,str)},
                'total_formula':total.value,'total_comment':total.comment.text if total.comment else None,
                'current_total':value(values.cell(bottom,qi).value),'totals_row_count':table.totalsRowCount,
                'hidden_transaction_rows':sum(bool(values.row_dimensions[r].hidden) for r in range(top+1,bottom)),
                'transaction_rows':bottom-top-1,
                'rows':[{'row':r,'date':value(values.cell(r,di).value),'quantity':value(values.cell(r,qi).value)} for r in range(top+1,min(bottom,top+65))]})
    return result


def inspect(content):
    current=extract_quantities(content)
    book=load_workbook(BytesIO(content),data_only=False,keep_links=False)
    cached=load_workbook(BytesIO(content),data_only=True,keep_links=False)
    try:
        spec=MAPPING['Tether']
        sheet,table=next((s,t) for s in book for t in s.tables.values() if t.name==spec['table'])
        left,top,right,bottom=range_boundaries(table.ref)
        headers=[c.name for c in table.tableColumns]
        qi=left+headers.index(spec['column']);di=left+headers.index('Date')
        total=sheet.cell(bottom,qi)
        first=top+1
        _,trades,problems,_=reconstruct(content,{})
        daily=defaultdict(float)
        for day,amount in trades['Tether']:daily[day]+=amount
        start=date(2025,10,25);balance=0.0;negative_before=False;negative_within=False
        for day in sorted(daily):
            balance+=daily[day]
            if balance < -max(1e-8,abs(current['Tether'])*1e-10):
                if day<start:negative_before=True
                else:negative_within=True
        opening=sum(q for d,q in trades['Tether'] if d<start)
        negative_within=negative_within or opening < -max(1e-8,abs(current['Tether'])*1e-10)
        retained=sum(cached[sheet.title].cell(r,qi).value for r in range(first,bottom)
            if isinstance(cached[sheet.title].cell(r,qi).value,(int,float)) and not isinstance(cached[sheet.title].cell(r,qi).value,bool)
            and norm(str(cached[sheet.title].cell(r,di).value or ''))!=norm('هزینه گم شده'))
        q=cached[sheet.title].cell(first,qi).value
        def value(v):return v.isoformat() if isinstance(v,(date,datetime)) else v
        return {'workbook_sha256':hashlib.sha256(content).hexdigest(),'sheet':sheet.title,'table':table.name,
            'total_cell':total.coordinate,'total_formula':total.value,'total_comment':total.comment.text if total.comment else None,
            'current_total':current['Tether'],'retained_dated_net':retained,
            'reconstructed_net':sum(q for _,q in trades['Tether']),
            'negative_before_analysis_window':negative_before,'negative_within_analysis_window':negative_within,
            'first_quantity_cell':sheet.cell(first,qi).coordinate,'first_quantity':q,'first_date':value(cached[sheet.title].cell(first,di).value),
            'first_row_text':{sheet.cell(top,c).value:str(cached[sheet.title].cell(first,c).value)[:500] for c in range(left,right+1) if isinstance(cached[sheet.title].cell(first,c).value,str)},
            'residual_after_first_row_exclusion':current['Tether']-(retained-q) if isinstance(q,(int,float)) else None,
            'ignored_rows':[{'cell':sheet.cell(r,qi).coordinate,'quantity':cached[sheet.title].cell(r,qi).value} for r in range(first,bottom) if norm(str(cached[sheet.title].cell(r,di).value or ''))==norm('هزینه گم شده')],
            'issues':problems['Tether'],'dust_and_unmapped_tables':inspect_dust_tables(book,cached)}
    finally:book.close();cached.close()


def run():
    certificate=base64.b64decode(required('DIAGNOSTIC_CERTIFICATE'),validate=True)
    if len(certificate)>8192 or b'BEGIN CERTIFICATE' not in certificate:raise SyncError('invalid_diagnostic_certificate')
    content,_=download_workbook(graph_token())
    payload=json.dumps(inspect(content),ensure_ascii=False).encode()
    if len(payload)>32768:raise SyncError('private_diagnostic_too_large')
    with tempfile.NamedTemporaryFile() as cert:
        cert.write(certificate);cert.flush()
        result=subprocess.run(['openssl','cms','-encrypt','-aes-256-cbc','-binary','-outform','DER',cert.name],input=payload,capture_output=True)
    if result.returncode:raise SyncError('private_diagnostic_encryption_failed')
    body={'name':'Private Tether total inspection','head_sha':required('GITHUB_SHA'),'status':'completed','conclusion':'success',
          'output':{'title':'Encrypted formula inspection for the requesting Codex session','summary':base64.b64encode(result.stdout).decode()}}
    env={**os.environ,'GH_TOKEN':required('PRIVATE_AUDIT_CHECK_TOKEN')}
    result=subprocess.run(['gh','api','--method','POST','repos/'+required('GITHUB_REPOSITORY')+'/check-runs','--input','-'],input=json.dumps(body),capture_output=True,text=True,env=env)
    if result.returncode:raise SyncError('private_diagnostic_publication_failed')
    return {'ok':True,'encrypted_diagnostic_published':True}


if __name__=='__main__':
    try:print(json.dumps(run()))
    except Exception as error:
        print(json.dumps({'ok':False,'error':str(error) if isinstance(error,SyncError) else 'private_diagnostic_failed'}))
        raise SystemExit(1)
