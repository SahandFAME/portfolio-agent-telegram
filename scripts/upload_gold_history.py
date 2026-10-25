"""Publish only the generated report into the owner's OneDrive اسناد folder.

Separate delegated credential; no source workbook writes or sharing operations.
Public output contains fixed status fields only, never drive links or contents.
"""
import base64
import hashlib
from io import BytesIO
import json
import os
import re
import subprocess
from urllib.parse import quote, urlencode, urlsplit
import zipfile

from openpyxl import load_workbook
from gold_history import api
from verify_onedrive import SyncError, http, json_request, required

GRAPH='https://graph.microsoft.com/v1.0'
DRIVE='daef687b3c624f3'
SOURCE_ITEM='DAEF687B3C624F3!2063'
FOLDER='اسناد'
FILENAME='Portfolio-18K-Gold-History.xlsx'
IDENTIFIER='portfolio-agent:18k-gold-history:v1'
MAX_REPORT_BYTES=3*1024*1024


def same_drive(value):
    # Personal OneDrive CID values may be returned with a leading zero.
    # Compare only bounded hexadecimal IDs, never arbitrary suffixes or names.
    return bool(isinstance(value,str) and re.fullmatch(r'[0-9a-fA-F]{1,16}',value)
                and int(value,16)==int(DRIVE,16))


def export_token():
    fields={'client_id':required('MS_GRAPH_CLIENT_ID'), 'grant_type':'refresh_token',
            'refresh_token':required('MS_GRAPH_EXPORT_REFRESH_TOKEN'),
            'scope':'https://graph.microsoft.com/Files.ReadWrite'}
    token=json_request('https://login.microsoftonline.com/consumers/oauth2/v2.0/token',
        method='POST',headers={'Content-Type':'application/x-www-form-urlencoded'},body=urlencode(fields).encode())
    scopes={s.rsplit('/',1)[-1].lower() for s in token.get('scope','').split()}
    if not token.get('access_token') or 'files.readwrite' not in scopes:
        raise SyncError('report_write_scope_missing')
    if token.get('refresh_token'):
        required('GH_TOKEN')
        if required('GITHUB_REPOSITORY')!='SahandFAME/portfolio-agent-telegram':
            raise SyncError('unexpected_repository')
        result=subprocess.run(['gh','secret','set','MS_GRAPH_EXPORT_REFRESH_TOKEN','--repo',os.environ['GITHUB_REPOSITORY']],
            input=token['refresh_token'].encode(),stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        if result.returncode:raise SyncError('export_refresh_token_persistence_failed')
    return token['access_token']


def validate_report(data):
    if not isinstance(data,bytes) or not 0<len(data)<=MAX_REPORT_BYTES:
        raise SyncError('invalid_report_size')
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist())>30*1024*1024 or any('vbaproject' in i.lower() for i in archive.namelist()):
                raise ValueError()
        book=load_workbook(BytesIO(data),read_only=True,keep_links=False)
        try:
            if book.properties.identifier!=IDENTIFIER or book.sheetnames!=['ارزش هفتگی','دارایی‌ها','ترکیب دارایی','موارد نیازمند بررسی','خلاصه تحلیل']:
                raise ValueError()
        finally:book.close()
    except Exception:
        raise SyncError('report_identity_not_verified') from None


def check_item(item,folder_id):
    parent=item.get('parentReference',{})
    if not item.get('id') or item['id'].lower()==SOURCE_ITEM.lower() or item.get('name')!=FILENAME or not isinstance(item.get('file'),dict) or not item.get('eTag') or parent.get('id')!=folder_id or not same_drive(parent.get('driveId')):
        raise SyncError('report_destination_not_verified')


def download_report(item,headers,folder_id):
    check_item(item,folder_id)
    meta=json_request(f'{GRAPH}/drives/{DRIVE}/items/{quote(item["id"],safe="")}',headers=headers)
    check_item(meta,folder_id)
    url=meta.get('@microsoft.graph.downloadUrl','')
    if urlsplit(url).scheme!='https':raise SyncError('report_download_unavailable')
    # Preauthenticated OneDrive URL receives no Graph Authorization header.
    data=http(url)
    validate_report(data)
    return data


def upload_report():
    report=api('?data=gold-history-file')
    try:data=base64.b64decode(report['content_base64'],validate=True)
    except Exception:raise SyncError('invalid_private_report') from None
    if hashlib.sha256(data).hexdigest()!=report.get('sha256'):raise SyncError('private_report_hash_mismatch')
    validate_report(data)
    token=export_token();headers={'Authorization':'Bearer '+token}
    drive=json_request(GRAPH+'/me/drive?$select=id',headers=headers)
    if not same_drive(drive.get('id')):raise SyncError('unexpected_onedrive_account')
    folder=json_request(f'{GRAPH}/drives/{DRIVE}/root:/{quote(FOLDER,safe="")}?$select=id,name,folder,parentReference',headers=headers)
    if folder.get('name')!=FOLDER or not isinstance(folder.get('folder'),dict) or not folder.get('id') or not same_drive(folder.get('parentReference',{}).get('driveId')):
        raise SyncError('report_folder_not_verified')
    folder_id=folder['id'];base=f'{GRAPH}/drives/{DRIVE}/items/{quote(folder_id,safe="")}'
    target=base+':/'+quote(FILENAME,safe='')
    try:existing=json_request(target,headers=headers)
    except SyncError as error:
        if str(error)!='http_404':raise
        existing=None
    if existing:
        check_item(existing,folder_id)
        # Refuse to replace an unrelated workbook with the same filename.
        previous=download_report(existing,headers,folder_id)
        target=f'{GRAPH}/drives/{DRIVE}/items/{quote(existing["id"],safe="")}/content'
        write_headers={**headers,'If-Match':existing['eTag']}
    else:
        previous=None
        target+=':/content?@microsoft.graph.conflictBehavior=fail'
        write_headers={**headers,'If-None-Match':'*'}
    # Avoid publishing an older revision after another admin changes history.
    if api('?data=gold-history').get('revision')!=report.get('revision'):
        raise SyncError('history_revision_changed_before_onedrive_upload')
    if previous!=data:
        uploaded=json_request(target,method='PUT',headers={**write_headers,'Content-Type':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'},body=data,retries=0)
        check_item(uploaded,folder_id)
        if existing and uploaded['id']!=existing['id']:raise SyncError('report_item_identity_changed')
    else:uploaded=existing
    if hashlib.sha256(download_report(uploaded,headers,folder_id)).hexdigest()!=report['sha256']:
        raise SyncError('onedrive_report_readback_mismatch')
    if api('?data=gold-history').get('revision')!=report.get('revision'):
        raise SyncError('history_revision_changed_after_onedrive_upload')
    return {'ok':True,'onedrive_uploaded':True,'onedrive_hash_verified':True}


if __name__=='__main__':
    try:print(json.dumps(upload_report()))
    except Exception as error:
        print(json.dumps({'ok':False,'error':str(error) if isinstance(error,SyncError) else 'onedrive_export_failed'}))
        raise SystemExit(1)
