import importlib.util
from io import BytesIO
import json
import os
from pathlib import Path
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
import zipfile

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableColumn

SPEC = importlib.util.spec_from_file_location('portfolio_sync', Path(__file__).parents[1] / 'scripts/portfolio_sync.py')
sync = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(sync)
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def workbook(*, formula_cache=True, no_totals=False, omit=False):
    """Entirely synthetic workbook; no private workbook or real balances."""
    book = Workbook()
    book.remove(book.active)
    for i, spec in enumerate(sync.MAPPING.values()):
        if omit and i == 0:
            continue
        sheet = book.create_sheet(f'Fixture{i}')
        sheet.append([spec['column']])
        sheet.append([999999])  # Deliberately misleading transaction value.
        sheet.append(['=SUBTOTAL(109,A2:A2)'])
        column = TableColumn(id=1, name=spec['column'], totalsRowFunction='sum')
        table = Table(displayName=spec['table'], ref='A1:A3',
                      totalsRowCount=0 if no_totals else 1, tableColumns=[column])
        sheet.add_table(table)
    stream = BytesIO()
    book.save(stream)  # Generated fixture only. Runtime never saves a workbook.
    out = BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(out, 'w') as dest:
        for name in source.namelist():
            data = source.read(name)
            if name.startswith('xl/worksheets/sheet') and formula_cache:
                root = ET.fromstring(data)
                root.find('.//s:c[@r="A3"]/s:v', NS).text = '0'
                data = ET.tostring(root)
            dest.writestr(name, data)
    return out.getvalue()


class ExtractionTests(unittest.TestCase):
    def test_explicit_cached_totals_not_transactions(self):
        content = workbook()
        quantities = sync.extract_quantities(content)
        self.assertEqual(len(quantities), 23)
        self.assertEqual(set(quantities.values()), {0})
        self.assertEqual(sync.extract_quantities(content), quantities)

    def test_missing_formula_cache_fails_closed(self):
        with self.assertRaisesRegex(sync.SyncError, 'quantity_must_be_numeric'):
            sync.extract_quantities(workbook(formula_cache=False))

    def test_no_explicit_totals_fails_closed(self):
        with self.assertRaisesRegex(sync.SyncError, 'current_balance_total_not_explicit'):
            sync.extract_quantities(workbook(no_totals=True))

    def test_missing_table_fails_closed(self):
        with self.assertRaisesRegex(sync.SyncError, 'required_table_missing'):
            sync.extract_quantities(workbook(omit=True))

    def test_invalid_archive(self):
        with self.assertRaisesRegex(sync.SyncError, 'invalid_workbook'):
            sync.extract_quantities(b'not a workbook')

    def test_strict_quantity_types(self):
        for value in [None, True, '', '2', -1, float('nan'), float('inf')]:
            with self.subTest(value=value):
                assets = dict.fromkeys(sync.MAPPING, 0)
                assets['BTC'] = value
                with self.assertRaises(sync.SyncError):
                    sync.validate_assets(assets)

    def test_manual_assets_cannot_be_submitted(self):
        assets = dict.fromkeys(sync.MAPPING, 0)
        assets['دلار'] = 3030
        with self.assertRaisesRegex(sync.SyncError, 'incorrect_asset_set'):
            sync.validate_assets(assets)


class IntegrationTests(unittest.TestCase):
    def setUp(self):
        self.assets = dict.fromkeys(sync.MAPPING, 0)
        self.stamp = '2026-01-01T00:00:00Z'
        self.snapshot = {'updated_at': self.stamp, 'workbook_updated_at': self.stamp,
                         'assets': self.assets}
        self.env = patch.dict(os.environ, {'PORTFOLIO_ENDPOINT': 'https://fixture.invalid/api/telegram',
                                         'PORTFOLIO_SYNC_SECRET_AGENT': 'test-fixture'}, clear=True)
        self.env.start()
        self.addCleanup(self.env.stop)

    def test_default_compare_never_posts_and_unchanged_file_is_valid(self):
        with patch.object(sync, 'read_snapshot', return_value=self.snapshot), patch.object(sync, 'json_request') as post:
            result = sync.synchronize(self.assets, self.stamp)
            self.assertTrue(result['matched'])
            self.assertFalse(result['production_written'])
            post.assert_not_called()

    def test_older_source_is_rejected(self):
        with patch.object(sync, 'read_snapshot', return_value=self.snapshot), self.assertRaisesRegex(sync.SyncError, 'older_workbook_rejected'):
            sync.synchronize(self.assets, '2025-01-01T00:00:00Z')

    def test_future_source_is_rejected(self):
        with self.assertRaisesRegex(sync.SyncError, 'future_workbook_timestamp'):
            sync.synchronize(self.assets, '2099-01-01T00:00:00Z')

    def test_write_requires_approval(self):
        with patch.object(sync, 'read_snapshot', return_value=self.snapshot), patch.object(sync, 'json_request') as post:
            with self.assertRaisesRegex(sync.SyncError, 'migration_not_approved'):
                sync.synchronize(self.assets, self.stamp, write=True)
            post.assert_not_called()

    def test_approved_write_requires_exact_authenticated_readback(self):
        with patch.dict(os.environ, {'PORTFOLIO_MIGRATION_APPROVED': 'true'}), \
             patch.object(sync, 'utcnow', return_value=self.stamp), \
             patch.object(sync, 'read_snapshot', return_value=self.snapshot) as read, \
             patch.object(sync, 'json_request', return_value={'ok':True, 'received':23, 'stored':True}) as post:
            result = sync.synchronize(self.assets, self.stamp, write=True)
            self.assertTrue(result['readback_verified'])
            self.assertEqual(read.call_count, 2)
            self.assertEqual(post.call_args.kwargs['retries'], 0)
            self.assertEqual(set(json.loads(post.call_args.kwargs['body'])['assets']), set(sync.MAPPING))

    def test_readback_mismatch_is_failure(self):
        with self.assertRaisesRegex(sync.SyncError, 'snapshot_readback_mismatch'):
            sync.verify_readback(self.snapshot, {**self.snapshot, 'updated_at':'2026-01-02T00:00:00Z'})

    def test_storage_read_failure_prevents_write(self):
        with patch.object(sync, 'read_snapshot', side_effect=sync.SyncError('http_500')), patch.object(sync, 'json_request') as post:
            with self.assertRaises(sync.SyncError):
                sync.synchronize(self.assets, self.stamp, write=True)
            post.assert_not_called()

    def test_stale_monitor_fails(self):
        with patch.object(sync, 'json_request', return_value={'ok':True,'service':'Portfolio Agent Telegram Bot'}), \
             patch.object(sync, 'read_snapshot', return_value=self.snapshot), self.assertRaisesRegex(sync.SyncError, 'snapshot_stale_or_future'):
            sync.monitor()

    def test_graph_bearer_not_forwarded_to_file_host(self):
        metadata = {'name':'Trading Journal.xlsm','file':{},'eTag':'fixture',
                    'lastModifiedDateTime':self.stamp}
        metadata['file'] = {'mimeType':'fixture'}
        with patch.dict(os.environ, {'MS_GRAPH_DRIVE_ID':'fixture','MS_GRAPH_ITEM_ID':'fixture'}), \
             patch.object(sync, 'json_request', side_effect=[metadata, {'@microsoft.graph.downloadUrl':'https://file.fixture.invalid/read'}, {'eTag':'fixture'}]), \
             patch.object(sync, 'http', return_value=b'fixture') as download:
            sync.download_workbook('fixture-token')
            self.assertNotIn('headers', download.call_args.kwargs)

    def test_rotated_refresh_token_persisted_via_stdin_without_logging(self):
        with patch.dict(os.environ, {'MS_GRAPH_CLIENT_ID':'fixture','MS_GRAPH_REFRESH_TOKEN':'old-fixture',
                                   'GH_TOKEN':'fixture', 'GITHUB_REPOSITORY':'fixture/repo'}), \
             patch.object(sync, 'json_request', return_value={'access_token':'access-fixture','refresh_token':'new-fixture'}), \
             patch.object(sync.subprocess, 'run') as run:
            run.return_value.returncode = 0
            self.assertEqual(sync.graph_token(), 'access-fixture')
            self.assertEqual(run.call_args.kwargs['input'], b'new-fixture')
            self.assertNotIn('new-fixture', ' '.join(run.call_args.args[0]))

    def test_refresh_persistence_failure_stops_sync(self):
        with patch.dict(os.environ, {'MS_GRAPH_CLIENT_ID':'fixture','MS_GRAPH_REFRESH_TOKEN':'old-fixture'}), \
             patch.object(sync, 'json_request', return_value={'access_token':'fixture','refresh_token':'fixture'}), \
             self.assertRaisesRegex(sync.SyncError, 'refresh_token_persistence_not_configured'):
            sync.graph_token()

    def test_readonly_telegram_inspection_does_not_send_messages(self):
        with patch.dict(os.environ, {'TELEGRAM_BOT_TOKEN':'fixture'}), \
             patch.object(sync, 'json_request', side_effect=[{'ok':True,'result':{'username':'SaRoPortfolioAgentBot'}},
                {'ok':True,'result':{'url':'https://fixture.invalid/api/telegram','pending_update_count':0}}]) as call:
            result = sync.inspect_telegram()
            self.assertFalse(result['telegram_commands_verified'])
            self.assertTrue(result['webhook_url_verified'])
            self.assertTrue(call.call_args_list[0].args[0].endswith('/getMe'))
            self.assertTrue(call.call_args_list[1].args[0].endswith('/getWebhookInfo'))


if __name__ == '__main__':
    unittest.main()
