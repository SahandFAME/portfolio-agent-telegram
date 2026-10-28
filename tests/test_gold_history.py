from datetime import date,datetime,timezone
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
import zipfile

from openpyxl import Workbook,load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table,TableColumn
sys.path.insert(0,str(Path(__file__).parents[1]/'scripts'))
import gold_history as history
from verify_onedrive import MAPPING


def source_workbook(*,ambiguous=False,mismatch=False,extra=False,date_overrides=None,tether_accounting=False,dust_ada=False,dust_assets=(),matic_remainder=False,matic_date=None):
    book=Workbook();book.remove(book.active);totals={}
    for i,(asset,spec) in enumerate(MAPPING.items()):
        sheet=book.create_sheet('Fixture'+str(i));sheet.append([spec['column'],'Date' if asset in history.CRYPTO else 'تاریخ'])
        if asset=='Tether' and tether_accounting:
            sheet.append([7,datetime(2023,1,1)]);sheet.append([20,datetime(2025,11,1)])
            sheet.append([-2,'هزینه گم شده']);sheet.append([-5,datetime(2026,1,1)]);total=10
        elif asset in dust_assets or asset=='ADA' and dust_ada:
            sheet.append([10,'08/02/1404' if asset=='ADA' else datetime(2025,4,28)]);sheet.append([-9.75,datetime(2026,9,26)]);total=0
        elif asset=='بلک راک':
            sheet.append([15,datetime(2026,9,28)]);sheet.append([None,None]);total=15
        else:
            sheet.append([10,'01/02/2025' if ambiguous and i==0 else datetime(2025,2,1)])
            sheet.append([-4,datetime(2026,9,26)]);total=6
        bottom=sheet.max_row+1
        sheet.append(['=SUBTOTAL(109,Table13444[Quantity])-A2-3' if asset=='Tether' and tether_accounting else '=SUBTOTAL(109,A2:A3)',None])
        if asset=='Tether' and tether_accounting:sheet.cell(bottom,1).comment=Comment('Sum minus "My Trading Capital" and ~3$ for gas','Synthetic')
        for row,value in (date_overrides or {}).get(asset,{}).items():sheet.cell(row,2).value=value
        cols=[TableColumn(id=1,name=spec['column'],totalsRowFunction='sum'),TableColumn(id=2,name=sheet.cell(1,2).value)]
        sheet.add_table(Table(displayName=spec['table'],ref=f'A1:B{bottom}',tableColumns=cols,totalsRowCount=1))
        totals[f'xl/worksheets/sheet{i+1}.xml']=(bottom,total+(1 if mismatch and i==0 else 0))
    if extra:
        sheet=book.create_sheet('OldAssetFixture');sheet.append(['Quantity','Date']);sheet.append([1,datetime(2025,1,1)]);sheet.append([0,None]);
        sheet.add_table(Table(displayName='UnknownOldHolding',ref='A1:B3',tableColumns=[TableColumn(id=1,name='Quantity'),TableColumn(id=2,name='Date')],totalsRowCount=1))
    if matic_remainder:
        sheet=book.create_sheet('Crypto Investing Portfolio');sheet.append(['Matic']);sheet.append(['Quantity','Date'])
        sheet.append([10,datetime(2024,1,1)]);sheet.append([-9.5,matic_date or datetime(2025,1,1)])
        sheet.append(['=SUBTOTAL(109,Table1312131921[Quantity])-0.5',None])
        sheet.add_table(Table(displayName='Table1312131921',ref='A2:B5',tableColumns=[TableColumn(id=1,name='Quantity'),TableColumn(id=2,name='Date')],totalsRowCount=1))
        totals[f'xl/worksheets/sheet{len(book.worksheets)}.xml']=(5,0)
    raw=BytesIO();book.save(raw);out=BytesIO();ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(raw) as src,zipfile.ZipFile(out,'w') as dst:
        for name in src.namelist():
            data=src.read(name)
            if name in totals:
                row,value=totals[name]
                root=ET.fromstring(data);root.find(f'.//s:c[@r="A{row}"]/s:v',ns).text=str(value);data=ET.tostring(root)
            dst.writestr(name,data)
    return out.getvalue()


class HistoricalQuantityTests(unittest.TestCase):
    def test_owner_waiver_ignores_only_the_identified_preperiod_matic_remainder(self):
        signature=history.hashlib.sha256(b'0.5').hexdigest()
        content=source_workbook(matic_remainder=True,extra=True);before=bytes(content)
        with patch.object(history,'MATIC_WAIVED_REMAINDER_SIGNATURE',signature):
            _,_,_,extras=history.reconstruct(content,{})
            self.assertEqual(extras,['OldAssetFixture / UnknownOldHolding'])
            _,_,_,extras=history.reconstruct(source_workbook(matic_remainder=True,matic_date='01/02/2025'),{})
            self.assertEqual(extras,[])
            for day in [datetime(2025,10,25),'not a date',datetime(2026,1,1)]:
                _,_,_,extras=history.reconstruct(source_workbook(matic_remainder=True,matic_date=day),{})
                self.assertIn('Crypto Investing Portfolio / Table1312131921',extras)
        _,_,_,extras=history.reconstruct(content,{})
        self.assertIn('Crypto Investing Portfolio / Table1312131921',extras)
        self.assertEqual(content,before)

    def test_matic_waiver_rejects_changed_identity_formula_amount_balance_or_hidden_rows(self):
        content=source_workbook(matic_remainder=True)
        book=load_workbook(BytesIO(content),data_only=True);formulas=load_workbook(BytesIO(content),data_only=False)
        s=book['Crypto Investing Portfolio'];f=formulas[s.title];table=s.tables['Table1312131921']
        events=[(date(2024,1,1),10),(date(2025,1,1),-9.5)]
        def accepted(balance=0):return history.owner_waived_matic_remainder(s,f,table,1,2,5,events,False,balance)
        try:
            with patch.object(history,'MATIC_WAIVED_REMAINDER_SIGNATURE',history.hashlib.sha256(b'0.5').hexdigest()):
                self.assertTrue(accepted())
                self.assertFalse(accepted(1));self.assertFalse(accepted(False))
                s['A1']='Other';self.assertFalse(accepted());s['A1']='Matic'
                s.row_dimensions[3].hidden=True;self.assertFalse(accepted());s.row_dimensions[3].hidden=False
                for value in ['=SUBTOTAL(109,Table1312131921[Quantity])-0.6','=UNTRUSTED()']:
                    f['A5']=value;self.assertFalse(accepted())
        finally:book.close();formulas.close()

    def test_verified_historical_crypto_prices_survive_provider_outage_without_bypassing_source_checks(self):
        dates=[datetime.combine(d,history.time(12),history.TEHRAN) for d in history.week_dates()]
        content=source_workbook(dust_assets={'ADA','SOL','ONDO','Link'})
        with patch.object(history,'tgju_history',return_value=[(d,1000000) for d in dates]),patch.object(history,'listed_history',return_value=[]),patch.object(history,'crypto_history',return_value=[(d,1) for d in dates]):
            rows,_=history.historical_weeks(content,'2026-10-24T00:00:00Z','2026-10-25T00:00:00Z',{'manual':{}},{})
        with patch.object(history,'tgju_history',return_value=[(d,1000000) for d in dates]),patch.object(history,'listed_history',return_value=[]),patch.object(history,'crypto_history',side_effect=TimeoutError()):
            recovered,diagnostics=history.historical_weeks(content,'2026-10-24T00:00:00Z','2026-10-25T00:00:00Z',{'manual':{}},{'weeks':rows})
        self.assertEqual(diagnostics['dust_candidate_assets'],4)
        self.assertEqual(diagnostics['dust_accepted_asset_weeks'],208)
        for old,new in zip(rows,recovered):
            for asset in ['ADA','SOL','ONDO','Link']:
                self.assertEqual(new['quotes'][asset],old['quotes'][asset])
                self.assertEqual(new['quantities'][asset],old['quantities'][asset])
        invalid=source_workbook(dust_assets={'ADA','SOL','ONDO','Link'},date_overrides={'SOL':{2:'not a date'}})
        with patch.object(history,'tgju_history',return_value=[(d,1000000) for d in dates]),patch.object(history,'listed_history',return_value=[]),patch.object(history,'crypto_history',side_effect=TimeoutError()):
            rejected,_=history.historical_weeks(invalid,'2026-10-24T00:00:00Z','2026-10-25T00:00:00Z',{'manual':{}},{'weeks':rows})
        self.assertTrue(all(r['quantities']['SOL'] is None for r in rejected))

    def test_crypto_evidence_cache_rejects_wrong_units_missing_fx_and_future_or_stale_quotes(self):
        stamp='2026-10-23T12:00:00+00:00'
        row={'mode':'historical','as_of':stamp,'quotes':{'SOL':{'price_rial':2000000,'quoted_at':stamp,'source':'CoinGecko USD × TGJU USD/IRR'}},
            'usd_quote':{'price_rial':1000000,'quoted_at':stamp,'source':'TGJU USD/IRR daily close'}}
        self.assertEqual(history.saved_crypto_series({'weeks':[row]},'SOL')[0][1],2)
        import copy
        changes=[lambda r:r.update(mode='live'),lambda r:r.update(usd_quote=None),
            lambda r:r['quotes']['SOL'].update(source='Unverified'),lambda r:r['usd_quote'].update(price_rial=-1),
            lambda r:r['quotes']['SOL'].update(quoted_at='2026-10-24T12:00:00+00:00'),
            lambda r:r['usd_quote'].update(quoted_at='2026-10-01T12:00:00+00:00')]
        for change in changes:
            bad=copy.deepcopy(row);change(bad)
            self.assertEqual(history.saved_crypto_series({'weeks':[bad]},'SOL'),[])

    def test_confirmed_sol_ondo_link_dust_uses_same_weekly_dollar_bound_without_waiving_other_assets(self):
        dates=[datetime.combine(d,history.time(12),history.TEHRAN) for d in history.week_dates()]
        content=source_workbook(dust_assets={'ADA','SOL','ONDO','Link','BTC'})
        with patch.object(history,'tgju_history',return_value=[(d,1000000) for d in dates]),patch.object(history,'listed_history',return_value=[]),patch.object(history,'crypto_history',return_value=[(d,1) for d in dates]):
            rows,_=history.historical_weeks(content,'2026-10-24T00:00:00Z','2026-10-25T00:00:00Z',{'manual':{}},{})
        for asset in ['ADA','SOL','ONDO','Link']:
            self.assertTrue(all(r['quantities'][asset] is not None for r in rows))
            self.assertEqual(rows[0]['quantities'][asset],10)
            self.assertEqual(rows[-1]['quantities'][asset],.25)
            self.assertTrue(all(not any(i.endswith(':'+asset) for i in r['issues']) for r in rows))
        self.assertTrue(all(r['quantities']['BTC'] is None for r in rows))

    def test_ada_dust_is_tolerated_per_week_only_with_valid_dated_usd_evidence(self):
        content=source_workbook(dust_ada=True);before=bytes(content)
        dates=[datetime.combine(d,history.time(12),history.TEHRAN) for d in history.week_dates()]
        def run(price,dates_override=None):
            raw=source_workbook(dust_ada=True,date_overrides=dates_override) if dates_override else content
            with patch.object(history,'tgju_history',return_value=[(d,1000000) for d in dates]),patch.object(history,'listed_history',return_value=[]),patch.object(history,'crypto_history',return_value=[(d,price) for d in dates]):
                return history.historical_weeks(raw,'2026-10-24T00:00:00Z','2026-10-25T00:00:00Z',{'manual':{}},{})[0]
        rows=run(1)
        self.assertEqual(rows[0]['quantities']['ADA'],10)  # Earlier material holdings survive.
        self.assertAlmostEqual(rows[-1]['quantities']['ADA'],.25)  # Raw quantity retained.
        self.assertTrue(all(not any('quantity_reconciliation_failed:ADA'==i for i in r['issues']) for r in rows))
        self.assertTrue(all(r['minimum_usd']==5 and r['usd_quote'] for r in rows))
        # Expensive or missing quotes cannot prove a sub-$5 discrepancy.
        self.assertIsNone(run(40)[-1]['quantities']['ADA'])
        self.assertIsNone(run(0)[-1]['quantities']['ADA'])
        self.assertIsNone(run(1,{'ADA':{2:'not a date'}})[-1]['quantities']['ADA'])
        self.assertIsNone(run(1,{'ADA':{2:'2026/10/01'}})[-1]['quantities']['ADA'])
        # Exactly $5 is outside the owner's strictly-less-than tolerance.
        self.assertIsNone(run(20)[-1]['quantities']['ADA'])
        self.assertEqual(content,before)

    def test_owner_confirmed_preperiod_deductions_adjust_every_week_without_a_fake_transaction(self):
        signature=history.deduction_signature(7,3,-2)
        with patch.object(history,'PREWINDOW_TETHER_SIGNATURE',signature):
            _,trades,problems,_=history.reconstruct(source_workbook(tether_accounting=True),{})
        # This synthetic fixture has no funds at the opening boundary, so the
        # adjustment correctly fails the nonnegative opening-balance check.
        self.assertIn('negative_reconstructed_balance',problems['Tether'])
        self.assertFalse(any(i.startswith('undated_') for i in problems['Tether']))
        events=history.HistoricalEvents([(date(2024,1,1),20),(date(2026,1,1),-5)])
        events.opening_adjustment=-5
        self.assertFalse(history.negative_historical_balance(events,10,start=date(2025,10,25)))
        self.assertEqual(history.quantities_at(date(2025,10,31),{'Tether':events},{'Tether':[]},{})['Tether'],15)
        self.assertEqual(history.quantities_at(date(2026,10,23),{'Tether':events},{'Tether':[]},{})['Tether'],10)
        self.assertEqual(len(events),2)
        self.assertNotEqual(history.deduction_signature(7,4,-2),signature)

    def test_reserved_capital_basis_validates_opening_and_in_window_balances(self):
        events=[(date(2023,1,1),-5),(date(2024,1,1),20),(date(2026,1,1),-1)]
        self.assertTrue(history.negative_historical_balance(events,14))
        self.assertFalse(history.negative_historical_balance(events,14,start=date(2025,10,25)))
        self.assertTrue(history.negative_historical_balance([(date(2024,1,1),-1),(date(2026,1,1),10)],9,start=date(2025,10,25)))
        self.assertTrue(history.negative_historical_balance([(date(2024,1,1),5),(date(2026,1,1),-6)],0,start=date(2025,10,25)))

    def test_documented_capital_and_fees_explain_current_total_without_inventing_expense_dates(self):
        content=source_workbook(tether_accounting=True);before=bytes(content)
        current,trades,problems,_=history.reconstruct(content,{})
        self.assertEqual(current['Tether'],10)
        self.assertEqual(sum(q for _,q in trades['Tether']),15)
        self.assertFalse(any(i.startswith('quantity_reconciliation_failed') for i in problems['Tether']))
        self.assertEqual(problems['Tether'],['undated_ignored_tether_adjustment@Table13444#quantity=-2','undated_tether_fee@A6#quantity=-3'])
        self.assertIsNone(history.quantities_at(date(2026,10,23),trades,problems,{})['Tether'])
        self.assertEqual(content,before)

    def test_total_rule_requires_exact_first_row_formula_matching_comment_and_unfiltered_rows(self):
        content=source_workbook(tether_accounting=True)
        formulas=load_workbook(BytesIO(content),data_only=False,keep_links=False)
        values=load_workbook(BytesIO(content),data_only=True,keep_links=False)
        try:
            sheet=next(s for s in formulas if 'Table13444' in s.tables);cached=values[sheet.title];table=sheet.tables['Table13444']
            self.assertEqual(history.tether_total_accounting(cached,sheet,table,1,2,1,6)['fee_quantity'],3)
            original=sheet['A6'].value
            for formula in ['=SUBTOTAL(109,Table13444[Quantity])-A3-3','=UNTRUSTED()-A2-3','=SUBTOTAL(109,Table13444[Quantity])-A2-4']:
                sheet['A6']=formula
                self.assertIsNone(history.tether_total_accounting(cached,sheet,table,1,2,1,6))
            sheet['A6']=original;cached.row_dimensions[3].hidden=True
            self.assertIsNone(history.tether_total_accounting(cached,sheet,table,1,2,1,6))
        finally:formulas.close();values.close()

    def test_confirmed_day_first_text_has_one_meaning_and_still_rejects_invalid_dates(self):
        for value,expected in [('01/02/2026',date(2026,2,1)),('۰۶/۰۷/۱۴۰۵',date(2026,9,28)),('2026/09/28',date(2026,9,28)),(datetime(2026,9,28),date(2026,9,28))]:
            self.assertEqual(history.parse_trade_date(value,day_first=True),expected)
        for value in ['09/28/2026','31/02/2026','01/02/26','2025/10/231','01/02/2026 & 02/02/2026']:
            with self.assertRaises(ValueError):history.parse_trade_date(value,day_first=True)
        with self.assertRaises(ValueError):history.parse_trade_date('01/02/2026')

    def test_day_first_confirmation_is_scoped_to_confirmed_assets_and_reconciles_weekly_quantities(self):
        content=source_workbook(date_overrides={a:{2:'01/02/2025',3:'02/03/2026'} for a in ['BTC','Link','Tether','ETH']})
        before=bytes(content)
        _,trades,problems,_=history.reconstruct(content,{})
        for asset in ['BTC','Link','Tether']:
            self.assertFalse(problems[asset]);self.assertEqual(trades[asset][0][0],date(2025,2,1))
            self.assertEqual(history.quantities_at(date(2026,2,20),trades,problems,{})[asset],10)
            self.assertEqual(history.quantities_at(date(2026,3,6),trades,problems,{})[asset],6)
        self.assertTrue(any(i.startswith('ambiguous_transaction_date') for i in problems['ETH']))
        self.assertEqual(content,before)

    def test_owner_ignored_tether_label_is_excluded_without_bypassing_balance_reconciliation(self):
        content=source_workbook(date_overrides={'Tether':{3:'هزینه گم شده'},'BTC':{3:'هزینه گم شده'}})
        before=bytes(content)
        current,trades,problems,_=history.reconstruct(content,{})
        self.assertEqual(current['Tether'],6)
        self.assertEqual(sum(q for _,q in trades['Tether']),10)
        self.assertEqual(problems['Tether'],['quantity_reconciliation_failed@Table13444#current_minus_expected_net=-4'])
        self.assertIsNone(history.quantities_at(date(2026,10,23),trades,problems,{})['Tether'])
        self.assertTrue(any(i.startswith('ambiguous_transaction_date') for i in problems['BTC']))
        self.assertEqual(content,before)

    def test_ignored_label_matches_only_the_exact_tether_date_note(self):
        for value,ignored in [('  هزینه\u200cگم شده  ',True),('هزینه گم شده دیگری',False),(None,False)]:
            _,_,problems,_=history.reconstruct(source_workbook(date_overrides={'Tether':{3:value}}),{})
            self.assertEqual(any(i.startswith('ambiguous_transaction_date') for i in problems['Tether']),not ignored)

    def test_dates_support_excel_solar_hijri_and_unambiguous_english_without_guessing(self):
        for value in [datetime(2026,9,28),'2026/09/28','1405/07/06','۱۴۰۵/۰۷/۰۶','09/28/2026','28/09/2026']:
            self.assertEqual(history.parse_trade_date(value),date(2026,9,28))
        self.assertEqual(history.parse_trade_date('1405/01/01'),date(2026,3,21))
        for value in [None,'01/02/2026','2025/10/231','2025/10/01 & 2025/10/02','هزینه گم شده']:
            with self.assertRaises(ValueError):history.parse_trade_date(value)

    def test_reconstruction_uses_signed_dated_trades_reconciled_to_current_totals_and_fund_launch(self):
        content=source_workbook();before=bytes(content)
        current,trades,problems,extras=history.reconstruct(content,{})
        self.assertFalse(extras);self.assertTrue(all(not issue for issue in problems.values()))
        first=history.quantities_at(date(2025,10,31),trades,problems,{'دلار':100})
        last=history.quantities_at(date(2026,10,23),trades,problems,{'دلار':100})
        self.assertEqual(first['طلا'],10);self.assertEqual(last['طلا'],6)
        self.assertEqual(first['بلک راک'],0);self.assertEqual(last['بلک راک'],15)
        self.assertEqual(last['دلار'],100);self.assertEqual(current['طلا'],last['طلا']);self.assertEqual(content,before)

    def test_ambiguous_dates_and_reconciliation_errors_remain_unknown_and_extra_trade_tables_are_detected(self):
        for kwargs in [{'ambiguous':True},{'mismatch':True}]:
            _,trades,problems,_=history.reconstruct(source_workbook(**kwargs),{})
            quantities=history.quantities_at(date(2026,10,23),trades,problems,{})
            self.assertIsNone(quantities['طلا']);self.assertTrue(problems['طلا'])
        _,_,_,extras=history.reconstruct(source_workbook(extra=True),{});self.assertEqual(extras,['OldAssetFixture / UnknownOldHolding'])

    def test_week_grid_has_52_actual_fridays_inside_the_confirmed_year(self):
        weeks=history.week_dates();self.assertEqual(len(weeks),52);self.assertEqual(len(set(weeks)),52)
        self.assertTrue(all(d.weekday()==4 for d in weeks));self.assertEqual(weeks[0],date(2025,10,31));self.assertEqual(weeks[-1],date(2026,10,23))

    def test_historical_quote_selection_never_uses_a_future_or_overly_stale_quote(self):
        cutoff=datetime(2026,10,23,23,59,tzinfo=history.TEHRAN)
        data=[(datetime(2026,10,1,tzinfo=history.TEHRAN),1),(datetime(2026,10,22,tzinfo=history.TEHRAN),2),(datetime(2026,10,24,tzinfo=history.TEHRAN),3)]
        self.assertEqual(history.last_quote(data,cutoff,'Synthetic')['price_rial'],2)
        self.assertIsNone(history.last_quote(data[:1],cutoff,'Synthetic'))

    def test_backfill_uses_historical_crypto_fx_units_and_does_not_price_fund_before_launch_or_use_guest_fallback(self):
        stamp=datetime(2026,10,20,tzinfo=history.TEHRAN)
        def tgju(slug):return [(stamp,200000000 if slug=='geram18' else 1000000)]
        config={'manual':{'سکه تمام':1,'ربع سکه بانکی':1,'ربع سکه غیره':1,'آبشده (طلب)':1,'آبشده (شمش زربد)':1,'دلار':100}}
        with patch.object(history,'tgju_history',side_effect=tgju),patch.object(history,'listed_history',return_value=[(stamp,100000)]),patch.object(history,'crypto_history',return_value=[(stamp,10)]):
            rows,_=history.historical_weeks(source_workbook(),'2026-10-24T00:00:00Z','2026-10-25T00:00:00Z',config,{})
        self.assertEqual(rows[-1]['quotes']['BTC']['price_rial'],10000000)
        self.assertEqual(rows[0]['quantities']['بلک راک'],0);self.assertNotIn('بلک راک',rows[-1]['quotes'])
        self.assertNotIn('شمش نقره 999',rows[-1]['quotes'])


class HistoryExcelTests(unittest.TestCase):
    def test_new_workbook_contains_real_chart_blank_gaps_sources_and_no_macros_or_formula_injection(self):
        def row(day,complete=True):
            return {'week_ending':day,'as_of':day+'T20:29:59Z','source_checked_at':'2026-10-25T10:00:00Z','mode':'historical',
                'status':'complete' if complete else 'incomplete','gold_grams':10 if complete else None,'total_rial':2000000000 if complete else None,
                'gold_quote':{'price_rial':200000000},'quantities':{'Stock':10,'Dust':.0001},'quotes':{
                    'Stock':{'price_rial':10000000,'source':'=HYPERLINK("unsafe")','quoted_at':day+'T00:00:00Z'},
                    'Dust':{'price_rial':1000,'source':'Synthetic','quoted_at':day+'T00:00:00Z'}},
                'allocation':{'Stock':100} if complete else None,'included_values':{'Stock':100000000} if complete else None,'missing_assets':[],'issues':[]}
        data=history.export_excel({'weeks':[row('2026-10-16'),row('2026-10-23',False)]},{'Stock':'سهام','Dust':'رمزارز'})
        book=load_workbook(BytesIO(data),data_only=False)
        self.assertEqual(book.sheetnames,['ارزش هفتگی','دارایی‌ها','ترکیب دارایی','موارد نیازمند بررسی','خلاصه تحلیل'])
        self.assertEqual(book.properties.identifier,'portfolio-agent:18k-gold-history:v1')
        self.assertEqual(book['ارزش هفتگی']['B2'].value,10);self.assertIsNone(book['ارزش هفتگی']['B3'].value)
        self.assertEqual(book['ارزش هفتگی']._charts[0].display_blanks,'gap');self.assertTrue(book['ارزش هفتگی']['I2'].font.bold)
        self.assertEqual(book['دارایی‌ها'].max_row,3);self.assertTrue(book['دارایی‌ها']['G2'].value.startswith("'="))
        # A verified asset remains useful when another holding blocks the total.
        self.assertEqual(book['دارایی‌ها']['F3'].value,10000000)
        self.assertTrue(book['دارایی‌ها']['A3'].font.bold);self.assertTrue(book['دارایی‌ها']['H3'].font.bold)
        self.assertIsNone(book['ارزش هفتگی']['E3'].value)
        self.assertEqual(book['ترکیب دارایی'].max_row,2)
        self.assertTrue(all(cell.data_type!='f' for sheet in book for rows in sheet for cell in rows))
        with zipfile.ZipFile(BytesIO(data)) as archive:self.assertFalse(any('vba' in name.lower() for name in archive.namelist()))
        book.close()

        filtered=row('2026-10-23',False);filtered['minimum_rial']=100000001
        book=load_workbook(BytesIO(history.export_excel({'weeks':[filtered]},{})))
        self.assertEqual(book['دارایی‌ها'].max_row,1);book.close()
        filtered['minimum_rial']=100000000
        book=load_workbook(BytesIO(history.export_excel({'weeks':[filtered]},{})))
        self.assertEqual(book['دارایی‌ها'].max_row,2);book.close()

class BoundedArchiveRequestTests(unittest.TestCase):
    def test_public_archive_request_has_a_hard_timeout_and_never_forwards_credentials(self):
        with patch.object(history.request,'build_opener') as factory:
            response=factory.return_value.open.return_value.__enter__.return_value
            response.read.return_value=b'{"data":[]}'
            self.assertEqual(history.market_json('https://api.tgju.org/v1/example'),{'data':[]})
            call=factory.return_value.open.call_args
            self.assertEqual(call.kwargs['timeout'],8)
            self.assertNotIn('Authorization',call.args[0].headers)
        with self.assertRaises(history.SyncError):history.market_json('https://example.invalid/private')

    def test_archive_timeout_produces_explicit_missing_quotes_without_losing_reconstructed_quantities(self):
        with patch.object(history,'market_json',side_effect=TimeoutError()):
            rows,_=history.historical_weeks(source_workbook(),'2026-10-24T00:00:00Z','2026-10-25T00:00:00Z',{'manual':{'دلار':100}}, {})
        self.assertEqual(len(rows),52);self.assertEqual(rows[-1]['quantities']['طلا'],6)
        self.assertIsNone(rows[-1]['gold_quote']);self.assertEqual(rows[-1]['quotes'],{})
