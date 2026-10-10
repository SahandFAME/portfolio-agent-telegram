import importlib.util
from io import BytesIO
import json
import os
from pathlib import Path
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
import zipfile

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableColumn

SPEC = importlib.util.spec_from_file_location('portfolio_sync', Path(__file__).parents[1] / 'scripts/verify_onedrive.py')
sync = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(sync)
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def workbook(*, formula_cache=True, no_totals=False, omit=False):
    """Entirely synthetic workbook; no private workbook or real balances."""
    book = Workbook()
    book.remove(book.active)
    for i, spec in enumerate(sync.MAPPING.values()):
        if omit and i == 0:
            continue
        sheet = book.create_sheet(f'Fixture{i}')
        sheet.append([spec['column']])
        sheet.append([999999])  # Deliberately misleading transaction value.
        sheet.append(['=SUBTOTAL(109,A2:A2)'])
        column = TableColumn(id=1, name=spec['column'], totalsRowFunction='sum')
        table = Table(displayName=spec['table'], ref='A1:A3',
                      totalsRowCount=0 if no_totals else 1, tableColumns=[column])
        sheet.add_table(table)
    stream = BytesIO()
    book.save(stream)  # Generated fixture only. Runtime never saves a workbook.
    out = BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(out, 'w') as dest:
        for name in source.namelist():
            data = source.read(name)
            if name.startswith('xl/worksheets/sheet') and formula_cache:
                root = ET.fromstring(data)
                root.find('.//s:c[@r="A3"]/s:v', NS).text = '0'
                data = ET.tostring(root)
            dest.writestr(name, data)
    return out.getvalue()


class ExtractionTests(unittest.TestCase):
    def test_explicit_cached_totals_not_transactions(self):
        content = workbook()
        quantities = sync.extract_quantities(content)
        self.assertEqual(len(quantities), 23)
        self.assertEqual(set(quantities.values()), {0})
        self.assertEqual(sync.extract_quantities(content), quantities)

    def test_missing_formula_cache_fails_closed(self):
        with self.assertRaisesRegex(sync.SyncError, 'quantity_must_be_numeric'):
            sync.extract_quantities(workbook(formula_cache=False))

    def test_no_explicit_totals_fails_closed(self):
        with self.assertRaisesRegex(sync.SyncError, 'current_balance_total_not_explicit'):
            sync.extract_quantities(workbook(no_totals=True))

    def test_missing_table_fails_closed(self):
        with self.assertRaisesRegex(sync.SyncError, 'required_table_missing'):
            sync.extract_quantities(workbook(omit=True))

    def test_invalid_archive(self):
        with self.assertRaisesRegex(sync.SyncError, 'invalid_workbook'):
            sync.extract_quantities(b'not a workbook')

    def test_strict_quantity_types(self):
        for value in [None, True, '', '2', -1, float('nan'), float('inf')]:
            with self.subTest(value=value):
                assets = dict.fromkeys(sync.MAPPING, 0)
                assets['BTC'] = value
                with self.assertRaises(sync.SyncError):
                    sync.validate_assets(assets)

    def test_manual_assets_cannot_be_submitted(self):
        assets = dict.fromkeys(sync.MAPPING, 0)
        assets['دلار'] = 3030
        with self.assertRaisesRegex(sync.SyncError, 'incorrect_asset_set'):
            sync.validate_assets(assets)



class GraphTests(unittest.TestCase):

    def setUp(self):
        self.assets = dict.fromkeys(sync.MAPPING, 0)
        self.stamp = '2026-01-01T00:00:00Z'
        self.snapshot = {'updated_at': self.stamp, 'workbook_updated_at': self.stamp,
                         'assets': self.assets}
        self.env = patch.dict(os.environ, {'PORTFOLIO_ENDPOINT': 'https://fixture.invalid/api/telegram',
                                         'PORTFOLIO_SYNC_SECRET_AGENT': 'test-fixture'}, clear=True)
        self.env.start()
        self.addCleanup(self.env.stop)

    def test_graph_bearer_not_forwarded_to_file_host(self):
        metadata = {'name':'Trading Journal.xlsm','file':{},'eTag':'fixture',
                    'lastModifiedDateTime':self.stamp}
        metadata['file'] = {'mimeType':'fixture'}
        with patch.dict(os.environ, {'MS_GRAPH_DRIVE_ID':'fixture','MS_GRAPH_ITEM_ID':'fixture'}), \
             patch.object(sync, 'json_request', side_effect=[metadata, {'@microsoft.graph.downloadUrl':'https://file.fixture.invalid/read'}, {'eTag':'fixture'}]), \
             patch.object(sync, 'http', return_value=b'fixture') as download:
            sync.download_workbook('fixture-token')
            self.assertNotIn('headers', download.call_args.kwargs)

    def test_rotated_refresh_token_persisted_via_stdin_without_logging(self):
        with patch.dict(os.environ, {'MS_GRAPH_CLIENT_ID':'fixture','MS_GRAPH_REFRESH_TOKEN':'old-fixture',
                                   'GH_TOKEN':'fixture', 'GITHUB_REPOSITORY':'fixture/repo'}), \
             patch.object(sync, 'json_request', return_value={'access_token':'access-fixture','refresh_token':'new-fixture'}), \
             patch.object(sync.subprocess, 'run') as run:
            run.return_value.returncode = 0
            self.assertEqual(sync.graph_token(), 'access-fixture')
            self.assertEqual(run.call_args.kwargs['input'], b'new-fixture')
            self.assertNotIn('new-fixture', ' '.join(run.call_args.args[0]))

    def test_refresh_persistence_failure_stops_sync(self):
        with patch.dict(os.environ, {'MS_GRAPH_CLIENT_ID':'fixture','MS_GRAPH_REFRESH_TOKEN':'old-fixture'}), \
             patch.object(sync, 'json_request', return_value={'access_token':'fixture','refresh_token':'fixture'}), \
             self.assertRaisesRegex(sync.SyncError, 'refresh_token_persistence_not_configured'):
            sync.graph_token()

    def test_write_permission_rejected_before_token_persistence(self):
        with patch.dict(os.environ, {'MS_GRAPH_CLIENT_ID':'fixture','MS_GRAPH_REFRESH_TOKEN':'fixture'}), \
             patch.object(sync, 'json_request', return_value={'access_token':'fixture','scope':'Files.ReadWrite'}), \
             patch.object(sync.subprocess, 'run') as store:
            with self.assertRaisesRegex(sync.SyncError, 'unexpected_write_scope'):
                sync.graph_token()
            store.assert_not_called()


class ComparisonTests(unittest.TestCase):
    def test_missing_secret_reports_precise_prerequisite(self):
        with patch.dict(os.environ, {}, clear=True), patch.object(sync, 'json_request') as network:
            self.assertEqual(sync.compare_production({})['production_comparison'], 'missing_repository_sync_secret')
            network.assert_not_called()

    def test_matching_and_different_cache_read_only(self):
        assets = dict.fromkeys(sync.MAPPING, 0)
        snapshot = {'updated_at': '2026-10-10T00:00:00Z', 'assets': dict(assets)}
        with patch.dict(os.environ, {'PORTFOLIO_SYNC_SECRET_AGENT':'fixture',
                                    'PORTFOLIO_ENDPOINT':'https://fixture.invalid/api/telegram'}, clear=True), \
             patch.object(sync, 'json_request', return_value={'ok':True, 'snapshot':snapshot}) as network:
            self.assertEqual(sync.compare_production(assets)['production_comparison'], 'matched')
            snapshot['assets']['BTC'] = 1
            self.assertEqual(sync.compare_production(assets)['changed_asset_count'], 1)
            for call in network.call_args_list:
                self.assertTrue(call.args[0].endswith('?data=portfolio'))
                self.assertNotIn('body', call.kwargs)
                self.assertNotIn('method', call.kwargs)

    def test_invalid_cache_is_not_reported_as_matching(self):
        with patch.dict(os.environ, {'PORTFOLIO_SYNC_SECRET_AGENT':'fixture',
                                    'PORTFOLIO_ENDPOINT':'https://fixture.invalid/api/telegram'}, clear=True), \
             patch.object(sync, 'json_request', return_value={'ok':True, 'snapshot':{'assets':{}}}):
            with self.assertRaisesRegex(sync.SyncError, 'incorrect_asset_set'):
                sync.compare_production(dict.fromkeys(sync.MAPPING, 0))


if __name__ == '__main__':
    unittest.main()
