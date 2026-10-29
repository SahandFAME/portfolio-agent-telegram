from datetime import date,datetime
from pathlib import Path
import sys
from unittest import TestCase
from unittest.mock import patch
from io import BytesIO
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).parents[1]/'scripts'))
import gold_history as h
from test_gold_history import source_workbook

class MonthlyHistoryTests(TestCase):
    def test_persian_month_end_grid_starts_inside_confirmed_manual_window_and_never_includes_open_month(self):
        days=h.month_dates(date(2026,10,29));self.assertEqual(len(days),12)
        self.assertEqual(days[0],date(2025,11,21));self.assertEqual(days[-1],date(2026,10,22))
        self.assertEqual(len(h.month_dates(date(2026,11,21))),12)
        self.assertEqual(h.month_dates(date(2026,11,22))[-1],date(2026,11,21))

    def test_month_end_holdings_are_reconstructed_from_trades_not_sampled_from_wednesday(self):
        content=source_workbook(additional=True,date_overrides={'طلا':{2:datetime(2025,11,20)}})
        dates=[datetime.combine(d,h.time(12),h.TEHRAN) for d in h.month_dates(date(2026,10,29))]
        with patch.object(h,'month_dates',return_value=[d.date() for d in dates]),patch.object(h,'tgju_history',return_value=[(d,1000000) for d in dates]),patch.object(h,'crypto_history',return_value=[(d,10) for d in dates]),patch.object(h,'listed_history',return_value=[(d,1000000) for d in dates]):
            rows,_=h.historical_weeks(content,'2026-10-24T00:00:00Z','2026-10-29T10:00:00Z',{'calendar':h.HISTORY_CALENDAR,'manual':{},'additional_assets':['XAUT','اهرم']},{})
        self.assertEqual(len(rows),12);self.assertTrue(all(r['period_kind']=='monthly' for r in rows))
        self.assertEqual(rows[0]['quantities']['طلا'],10);self.assertEqual(rows[-1]['quantities']['طلا'],6)
        self.assertEqual(rows[-2]['quantities']['بلک راک'],0);self.assertEqual(rows[-1]['quantities']['بلک راک'],15)

    def test_monthly_workbook_has_month_changes_and_marks_the_current_month_provisional(self):
        row={'period_kind':'monthly','mode':'baseline','week_ending':'2026-11-21','as_of':'2026-10-29T12:00:00Z','source_checked_at':'2026-10-29T12:00:00Z','quantities':{},'quotes':{},'status':'complete','total_rial':1000000,'gold_grams':1,'gold_quote':{'price_rial':1000000,'quoted_at':'2026-10-22T12:00:00Z'},'issues':[],'missing_assets':[],'valuation_basis':'estimated_last_known'}
        book=load_workbook(BytesIO(h.export_excel({'calendar':h.HISTORY_CALENDAR,'weeks':[row]},{})))
        self.assertEqual(book.sheetnames[0],'ارزش ماهانه');s=book.worksheets[0]
        self.assertEqual(s['C1'].value,'تغییر ماهانه (گرم)');self.assertIn('موقت',s['A2'].value);self.assertIn('برآوردی',s['G2'].value)
        self.assertTrue(s['I2'].font.bold);self.assertEqual(s._charts[0].display_blanks,'gap');book.close()

    def test_non_month_end_schedule_skips_without_reading_graph_or_exporting(self):
        class Fixed(datetime):
            @classmethod
            def now(cls,tz=None):return datetime(2026,10,29,12,tzinfo=h.TEHRAN)
        with patch.object(h,'datetime',Fixed),patch.object(h,'graph_token') as graph,patch.object(h,'publish_export') as export:
            result=h.run('capture');self.assertTrue(result['skipped']);graph.assert_not_called();export.assert_not_called()
